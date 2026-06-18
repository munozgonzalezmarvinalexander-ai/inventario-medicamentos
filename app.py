# -*- coding: utf-8 -*-
"""
Sistema de Inventario de Medicamentos
Backend Flask que usa archivos Excel como base de datos.

Hojas del archivo Excel:
- Inventario : stock actual de cada medicamento
- Entradas   : historial de medicamentos que entran
- Salidas    : historial de medicamentos entregados
- Reportes   : configuracion (umbral de bajo stock, etc.)
"""

import os
import io
from datetime import datetime
from threading import Lock

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


# ---------------------------------------------------------------------------
# Configuracion
# ---------------------------------------------------------------------------

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "data")
EXCEL_FILE = os.path.join(DATA_DIR, "inventario.xlsx")

# Umbral por defecto para alertas de bajo stock
DEFAULT_STOCK_BAJO = 10

# Evita que dos peticiones escriban el Excel al mismo tiempo
excel_lock = Lock()

app = Flask(__name__, static_folder="static", template_folder="templates")


# ---------------------------------------------------------------------------
# Utilidades de Excel
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
HEADER_FILL = PatternFill("solid", start_color="0F766E")  # teal-700
CENTER = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _style_header(ws, columnas):
    """Aplica estilo a la fila de encabezado."""
    for col_index, _ in enumerate(columnas, start=1):
        cell = ws.cell(row=1, column=col_index)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    # Anchos razonables
    anchos = [22, 16, 14, 22, 30]
    for i, ancho in enumerate(anchos[: len(columnas)], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho


def crear_excel_si_no_existe():
    """Crea el archivo Excel con todas las hojas si aun no existe."""
    if os.path.exists(EXCEL_FILE):
        return

    os.makedirs(DATA_DIR, exist_ok=True)
    wb = Workbook()

    # Hoja Inventario
    ws_inv = wb.active
    ws_inv.title = "Inventario"
    cols_inv = ["Nombre", "Cantidad", "Stock Minimo", "Ultima Actualizacion", "Observaciones"]
    ws_inv.append(cols_inv)
    _style_header(ws_inv, cols_inv)

    # Hoja Entradas
    ws_ent = wb.create_sheet("Entradas")
    cols_ent = ["Fecha", "Medicamento", "Cantidad", "Observaciones"]
    ws_ent.append(cols_ent)
    _style_header(ws_ent, cols_ent)

    # Hoja Salidas
    ws_sal = wb.create_sheet("Salidas")
    cols_sal = ["Fecha", "Medicamento", "Cantidad", "Recibido por", "Observaciones"]
    ws_sal.append(cols_sal)
    _style_header(ws_sal, cols_sal)

    # Hoja Reportes (configuracion)
    ws_rep = wb.create_sheet("Reportes")
    ws_rep.append(["Parametro", "Valor"])
    ws_rep.append(["Stock bajo (umbral global)", DEFAULT_STOCK_BAJO])
    _style_header(ws_rep, ["Parametro", "Valor"])

    wb.save(EXCEL_FILE)


def cargar_wb():
    """Carga el workbook (crea el archivo si no existe)."""
    crear_excel_si_no_existe()
    return load_workbook(EXCEL_FILE)


def obtener_umbral_bajo_stock(wb):
    """Devuelve el umbral global de bajo stock."""
    try:
        ws = wb["Reportes"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] and "stock bajo" in str(row[0]).lower():
                return int(row[1]) if row[1] is not None else DEFAULT_STOCK_BAJO
    except Exception:
        pass
    return DEFAULT_STOCK_BAJO


def leer_inventario(wb):
    """Devuelve la lista de medicamentos del inventario."""
    ws = wb["Inventario"]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            nombre, cantidad, stock_min, ultima, obs = (list(row) + [None] * 5)[:5]
            items.append({
                "nombre": str(nombre).strip(),
                "cantidad": int(cantidad) if cantidad is not None else 0,
                "stock_minimo": int(stock_min) if stock_min is not None else 0,
                "ultima_actualizacion": str(ultima) if ultima else "",
                "observaciones": str(obs) if obs else "",
            })
    return items


def buscar_fila_medicamento(ws_inv, nombre):
    """Devuelve el numero de fila del medicamento o None."""
    nombre_norm = nombre.strip().lower()
    for idx, row in enumerate(ws_inv.iter_rows(min_row=2, values_only=True), start=2):
        if row and row[0] and str(row[0]).strip().lower() == nombre_norm:
            return idx
    return None


def fecha_hoy():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


# ---------------------------------------------------------------------------
# Rutas: paginas
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template("index.html")


# ---------------------------------------------------------------------------
# Rutas API: Inventario
# ---------------------------------------------------------------------------

@app.route("/api/inventario", methods=["GET"])
def api_inventario():
    """Devuelve todo el inventario y el umbral de bajo stock."""
    with excel_lock:
        wb = cargar_wb()
        items = leer_inventario(wb)
        umbral = obtener_umbral_bajo_stock(wb)

    # Calcular alertas: usa stock_minimo si esta definido y > 0, si no usa el umbral global
    for it in items:
        umbral_efectivo = it["stock_minimo"] if it["stock_minimo"] > 0 else umbral
        it["umbral"] = umbral_efectivo
        it["bajo_stock"] = it["cantidad"] <= umbral_efectivo

    return jsonify({"inventario": items, "umbral_global": umbral})


@app.route("/api/dashboard", methods=["GET"])
def api_dashboard():
    """Devuelve estadisticas para el dashboard principal."""
    with excel_lock:
        wb = cargar_wb()
        items = leer_inventario(wb)
        umbral = obtener_umbral_bajo_stock(wb)

        # Conteos
        total_medicamentos = len(items)
        unidades_totales = sum(i["cantidad"] for i in items)
        bajo_stock = sum(
            1 for i in items
            if i["cantidad"] <= (i["stock_minimo"] if i["stock_minimo"] > 0 else umbral)
        )
        agotados = sum(1 for i in items if i["cantidad"] <= 0)

        # Entradas/salidas del mes en curso
        mes_actual = datetime.now().strftime("%Y-%m")
        entradas_mes = 0
        salidas_mes = 0

        ws_ent = wb["Entradas"]
        for row in ws_ent.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                if str(row[0]).startswith(mes_actual):
                    entradas_mes += int(row[2]) if row[2] else 0

        ws_sal = wb["Salidas"]
        for row in ws_sal.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                if str(row[0]).startswith(mes_actual):
                    salidas_mes += int(row[2]) if row[2] else 0

        # Top 5 mas entregados (historico)
        conteo_salidas = {}
        for row in ws_sal.iter_rows(min_row=2, values_only=True):
            if row and row[1] and row[2]:
                nombre = str(row[1]).strip()
                conteo_salidas[nombre] = conteo_salidas.get(nombre, 0) + int(row[2])
        top_entregados = sorted(conteo_salidas.items(), key=lambda x: x[1], reverse=True)[:5]
        top_entregados = [{"nombre": n, "total": c} for n, c in top_entregados]

        # Alertas de bajo stock (top 5 con menos stock)
        alertas = sorted(
            [i for i in items if i["cantidad"] <= (i["stock_minimo"] if i["stock_minimo"] > 0 else umbral)],
            key=lambda x: x["cantidad"]
        )[:5]

    return jsonify({
        "total_medicamentos": total_medicamentos,
        "unidades_totales": unidades_totales,
        "bajo_stock": bajo_stock,
        "agotados": agotados,
        "entradas_mes": entradas_mes,
        "salidas_mes": salidas_mes,
        "top_entregados": top_entregados,
        "alertas": alertas,
    })


# ---------------------------------------------------------------------------
# Rutas API: Entradas (agregar medicamentos)
# ---------------------------------------------------------------------------

@app.route("/api/entradas", methods=["POST"])
def api_agregar_entrada():
    """Registra entrada de medicamento. Si no existe en inventario, lo crea."""
    data = request.get_json(silent=True) or {}
    nombre = (data.get("nombre") or "").strip()
    cantidad = data.get("cantidad")
    stock_minimo = data.get("stock_minimo") or 0
    observaciones = (data.get("observaciones") or "").strip()
    fecha = (data.get("fecha") or "").strip() or fecha_hoy()

    if not nombre:
        return jsonify({"ok": False, "error": "El nombre del medicamento es obligatorio."}), 400
    try:
        cantidad = int(cantidad)
        stock_minimo = int(stock_minimo)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Cantidad invalida."}), 400
    if cantidad <= 0:
        return jsonify({"ok": False, "error": "La cantidad debe ser mayor que cero."}), 400

    with excel_lock:
        wb = cargar_wb()
        ws_inv = wb["Inventario"]
        ws_ent = wb["Entradas"]

        # Actualizar o crear en inventario
        fila = buscar_fila_medicamento(ws_inv, nombre)
        if fila is None:
            ws_inv.append([nombre, cantidad, stock_minimo, fecha, observaciones])
            nueva_fila = ws_inv.max_row
            for c in range(1, 6):
                ws_inv.cell(row=nueva_fila, column=c).border = THIN_BORDER
        else:
            actual = ws_inv.cell(row=fila, column=2).value or 0
            ws_inv.cell(row=fila, column=2).value = int(actual) + cantidad
            # Si proporciono nuevo stock minimo, actualizarlo
            if stock_minimo > 0:
                ws_inv.cell(row=fila, column=3).value = stock_minimo
            ws_inv.cell(row=fila, column=4).value = fecha
            if observaciones:
                ws_inv.cell(row=fila, column=5).value = observaciones

        # Registrar en historial de entradas
        ws_ent.append([fecha, nombre, cantidad, observaciones])
        nueva_fila_ent = ws_ent.max_row
        for c in range(1, 5):
            ws_ent.cell(row=nueva_fila_ent, column=c).border = THIN_BORDER

        wb.save(EXCEL_FILE)

    return jsonify({"ok": True, "mensaje": f"Entrada registrada: {cantidad} unidades de {nombre}."})


@app.route("/api/entradas", methods=["GET"])
def api_listar_entradas():
    """Lista todas las entradas, opcionalmente filtradas por fechas."""
    desde = request.args.get("desde", "")
    hasta = request.args.get("hasta", "")

    with excel_lock:
        wb = cargar_wb()
        ws = wb["Entradas"]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                fecha = str(row[0])[:10]
                if desde and fecha < desde:
                    continue
                if hasta and fecha > hasta:
                    continue
                rows.append({
                    "fecha": str(row[0]),
                    "medicamento": str(row[1]) if row[1] else "",
                    "cantidad": int(row[2]) if row[2] else 0,
                    "observaciones": str(row[3]) if row[3] else "",
                })

    rows.reverse()  # mas recientes primero
    return jsonify({"entradas": rows})


# ---------------------------------------------------------------------------
# Rutas API: Salidas (entregar medicamentos)
# ---------------------------------------------------------------------------

@app.route("/api/salidas", methods=["POST"])
def api_registrar_salida():
    """Descuenta del inventario y registra la salida."""
    data = request.get_json(silent=True) or {}
    nombre = (data.get("nombre") or "").strip()
    cantidad = data.get("cantidad")
    recibido_por = (data.get("recibido_por") or "").strip()
    observaciones = (data.get("observaciones") or "").strip()
    fecha = (data.get("fecha") or "").strip() or fecha_hoy()

    if not nombre:
        return jsonify({"ok": False, "error": "Selecciona un medicamento."}), 400
    if not recibido_por:
        return jsonify({"ok": False, "error": "Indica quien recibe el medicamento."}), 400
    try:
        cantidad = int(cantidad)
    except (TypeError, ValueError):
        return jsonify({"ok": False, "error": "Cantidad invalida."}), 400
    if cantidad <= 0:
        return jsonify({"ok": False, "error": "La cantidad debe ser mayor que cero."}), 400

    with excel_lock:
        wb = cargar_wb()
        ws_inv = wb["Inventario"]
        ws_sal = wb["Salidas"]

        fila = buscar_fila_medicamento(ws_inv, nombre)
        if fila is None:
            return jsonify({"ok": False, "error": f"'{nombre}' no existe en el inventario."}), 400

        actual = int(ws_inv.cell(row=fila, column=2).value or 0)
        if cantidad > actual:
            return jsonify({
                "ok": False,
                "error": f"Stock insuficiente. Disponible: {actual}, solicitado: {cantidad}."
            }), 400

        ws_inv.cell(row=fila, column=2).value = actual - cantidad
        ws_inv.cell(row=fila, column=4).value = fecha

        ws_sal.append([fecha, nombre, cantidad, recibido_por, observaciones])
        nueva_fila = ws_sal.max_row
        for c in range(1, 6):
            ws_sal.cell(row=nueva_fila, column=c).border = THIN_BORDER

        wb.save(EXCEL_FILE)

    return jsonify({
        "ok": True,
        "mensaje": f"Salida registrada: {cantidad} unidades de {nombre} para {recibido_por}.",
        "stock_restante": actual - cantidad,
    })


@app.route("/api/salidas", methods=["GET"])
def api_listar_salidas():
    """Lista las salidas con filtros opcionales."""
    desde = request.args.get("desde", "")
    hasta = request.args.get("hasta", "")
    medicamento = (request.args.get("medicamento") or "").strip().lower()

    with excel_lock:
        wb = cargar_wb()
        ws = wb["Salidas"]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                fecha = str(row[0])[:10]
                med = str(row[1]) if row[1] else ""
                if desde and fecha < desde:
                    continue
                if hasta and fecha > hasta:
                    continue
                if medicamento and medicamento not in med.lower():
                    continue
                rows.append({
                    "fecha": str(row[0]),
                    "medicamento": med,
                    "cantidad": int(row[2]) if row[2] else 0,
                    "recibido_por": str(row[3]) if row[3] else "",
                    "observaciones": str(row[4]) if row[4] else "",
                })

    rows.reverse()
    return jsonify({"salidas": rows})


# ---------------------------------------------------------------------------
# Rutas API: Reportes
# ---------------------------------------------------------------------------

@app.route("/api/reportes/por-mes", methods=["GET"])
def api_reporte_por_mes():
    """Devuelve totales de salidas agrupadas por mes (YYYY-MM)."""
    with excel_lock:
        wb = cargar_wb()
        ws = wb["Salidas"]
        por_mes = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[2]:
                mes = str(row[0])[:7]  # YYYY-MM
                por_mes[mes] = por_mes.get(mes, 0) + int(row[2])

    data = sorted(por_mes.items())
    return jsonify({"por_mes": [{"mes": m, "total": t} for m, t in data]})


@app.route("/api/reportes/top-entregados", methods=["GET"])
def api_reporte_top():
    """Devuelve los medicamentos mas entregados."""
    desde = request.args.get("desde", "")
    hasta = request.args.get("hasta", "")

    with excel_lock:
        wb = cargar_wb()
        ws = wb["Salidas"]
        conteo = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[1] and row[2]:
                fecha = str(row[0])[:10]
                if desde and fecha < desde:
                    continue
                if hasta and fecha > hasta:
                    continue
                nombre = str(row[1]).strip()
                conteo[nombre] = conteo.get(nombre, 0) + int(row[2])

    top = sorted(conteo.items(), key=lambda x: x[1], reverse=True)
    return jsonify({"top": [{"nombre": n, "total": t} for n, t in top]})


# ---------------------------------------------------------------------------
# Exportar Excel completo
# ---------------------------------------------------------------------------

@app.route("/api/exportar", methods=["GET"])
def api_exportar():
    """Devuelve el archivo Excel completo para descargar."""
    with excel_lock:
        crear_excel_si_no_existe()
        # Leer en memoria y enviar
        with open(EXCEL_FILE, "rb") as f:
            datos = f.read()

    nombre = f"inventario_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        io.BytesIO(datos),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nombre,
    )


@app.route("/api/exportar/reporte", methods=["GET"])
def api_exportar_reporte():
    """Genera un Excel con un reporte filtrado de salidas."""
    desde = request.args.get("desde", "")
    hasta = request.args.get("hasta", "")
    medicamento = (request.args.get("medicamento") or "").strip().lower()

    wb_out = Workbook()
    ws = wb_out.active
    ws.title = "Reporte de Salidas"
    columnas = ["Fecha", "Medicamento", "Cantidad", "Recibido por", "Observaciones"]
    ws.append(columnas)
    _style_header(ws, columnas)

    with excel_lock:
        wb = cargar_wb()
        ws_sal = wb["Salidas"]
        for row in ws_sal.iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                fecha = str(row[0])[:10]
                med = str(row[1]) if row[1] else ""
                if desde and fecha < desde:
                    continue
                if hasta and fecha > hasta:
                    continue
                if medicamento and medicamento not in med.lower():
                    continue
                ws.append([
                    str(row[0]),
                    med,
                    int(row[2]) if row[2] else 0,
                    str(row[3]) if row[3] else "",
                    str(row[4]) if row[4] else "",
                ])

    # Totalizador
    if ws.max_row > 1:
        fila_total = ws.max_row + 2
        ws.cell(row=fila_total, column=1, value="TOTAL UNIDADES").font = Font(bold=True)
        ws.cell(row=fila_total, column=3, value=f"=SUM(C2:C{ws.max_row})").font = Font(bold=True)

    buffer = io.BytesIO()
    wb_out.save(buffer)
    buffer.seek(0)

    nombre = f"reporte_salidas_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nombre,
    )


# ---------------------------------------------------------------------------
# Eliminar medicamento del inventario (opcional pero util)
# ---------------------------------------------------------------------------

@app.route("/api/inventario/<nombre>", methods=["DELETE"])
def api_eliminar_medicamento(nombre):
    """Elimina un medicamento del inventario (no borra el historial)."""
    with excel_lock:
        wb = cargar_wb()
        ws = wb["Inventario"]
        fila = buscar_fila_medicamento(ws, nombre)
        if fila is None:
            return jsonify({"ok": False, "error": "Medicamento no encontrado."}), 404
        ws.delete_rows(fila, 1)
        wb.save(EXCEL_FILE)
    return jsonify({"ok": True})


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    crear_excel_si_no_existe()
    print("=" * 60)
    print("  Sistema de Inventario de Medicamentos")
    print("=" * 60)
    print(f"  Archivo Excel: {EXCEL_FILE}")
    print(f"  Abre en tu navegador: http://localhost:5000")
    print("  Para detener: Ctrl + C")
    print("=" * 60)
    app.run(host="127.0.0.1", port=5000, debug=False)
